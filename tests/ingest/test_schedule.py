"""Excel schedule template and reader. Reader output is machine-read: provenance 'extracted', source 'xlsx'."""
import io
import zipfile
from pathlib import Path

import openpyxl
import pytest
from mep.engine.loader import load_pack
from mep.ingest.schedule import (
    MAX_ROWS,
    CellError,
    ScheduleInput,
    build_template,
    normalise_value,
    read_schedule,
    schedule_inputs,
)

ROOT = Path(__file__).resolve().parents[2]
EDITIONS = ("NCC2022", "NCC2025")


@pytest.fixture(scope="module")
def pack():
    return load_pack(ROOT / "rules")


def workbook(headers, rows, extra=None):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Systems"
    ws.append(headers)
    for r in rows:
        ws.append(r)
    if extra:
        extra(wb)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


H = ["system tag", "system type", "supply_airflow [L/s]", "control_deadband [K]"]
AC = "air_conditioning"


def read(data, pack, edition="NCC2022"):
    return read_schedule(data, pack, edition)


def test_template_columns_and_units(pack):
    wb = openpyxl.load_workbook(io.BytesIO(build_template(pack, "NCC2022")))
    assert wb.sheetnames[:3] == ["Systems", "ReadMe", "Lists"]
    header = [c.value for c in wb["Systems"][1]]
    assert header[:2] == ["system tag", "system type"]
    assert "control_deadband [K]" in header and "supply_airflow [L/s]" in header
    assert len(header) == 2 + len(schedule_inputs(pack, "NCC2022"))
    text = " ".join(str(c.value) for row in wb["ReadMe"].iter_rows() for c in row if c.value)
    assert "NCC2022" in text and "control_deadband" in text
    assert wb["Systems"].data_validations.dataValidation  # enumerations and bools are lists


def test_editions_are_separate(pack):
    a = set(schedule_inputs(pack, "NCC2022"))
    b = set(schedule_inputs(pack, "NCC2025"))
    assert a != b
    assert "max_airside_component_airflow" in b and "max_airside_component_airflow" not in a


@pytest.mark.parametrize("edition", EDITIONS)
def test_template_round_trips(pack, edition):
    result = read(build_template(pack, edition), pack, edition)
    assert result.errors == [] and result.rows == []


def _cells(data):
    wb = openpyxl.load_workbook(io.BytesIO(data) if isinstance(data, bytes) else data)
    return {ws.title: [[c.value for c in r] for r in ws.iter_rows()] for ws in wb}


def test_committed_templates_are_reproducible(pack):
    for edition in EDITIONS:
        committed = ROOT / "docs" / "templates" / f"mep-system-schedule-{edition}.xlsx"
        assert committed.exists(), "run scripts/build_schedule_template.py"
        assert _cells(committed) == _cells(build_template(pack, edition))


def test_good_row_extracted_provenance(pack):
    r = read(workbook(H, [["AHU-1", AC, 1200, 2]]), pack)
    assert r.errors == [] and len(r.rows) == 1
    row = r.rows[0]
    assert (row.tag, row.system_type) == ("AHU-1", AC)
    assert ScheduleInput("supply_airflow", 1200.0, "L/s", "extracted", "xlsx") in row.inputs
    assert all(i.provenance == "extracted" and i.source == "xlsx" for i in row.inputs)


def test_compatible_unit_converted_to_declared(pack):
    h = ["system tag", "system type", "supply_airflow [m^3/h]"]
    r = read(workbook(h, [["AHU-1", AC, 3600]]), pack)
    assert r.errors == []
    (i,) = r.rows[0].inputs
    assert i.unit == "L/s" and i.value == pytest.approx(1000.0)


def test_wrong_dimension_rejects_row(pack):
    h = ["system tag", "system type", "supply_airflow [kW]"]
    r = read(workbook(h, [["AHU-1", AC, 5], ["AHU-2", AC, None]]), pack)
    assert [x.tag for x in r.rows] == ["AHU-2"]
    assert [e.row for e in r.errors] == [2] and "supply_airflow" in r.errors[0].message


def test_absolute_temperature_not_mixed_with_difference(pack):
    h = ["system tag", "system type", "control_deadband [degC]"]
    r = read(workbook(h, [["AHU-1", AC, 2]]), pack)
    assert r.rows == [] and r.errors


def test_missing_unit_in_header(pack):
    h = ["system tag", "system type", "supply_airflow"]
    r = read(workbook(h, [["AHU-1", AC, 5]]), pack)
    assert r.rows == [] and "unit" in r.errors[0].message


def test_text_in_number_cell(pack):
    r = read(workbook(H, [["AHU-1", AC, "1200 L/s", 2]]), pack)
    assert r.rows == [] and r.errors[0].column == "supply_airflow"


def test_bool_and_nan_not_numbers(pack):
    r = read(workbook(H, [["A", AC, True, 2]]), pack)
    assert r.rows == [] and len(r.errors) == 1
    spec = schedule_inputs(pack, "NCC2022")["supply_airflow"]
    for bad in (float("nan"), float("inf"), True, "5"):
        with pytest.raises(CellError):
            normalise_value(spec, bad, "L/s")


def test_unknown_column_reported_row_kept(pack):
    r = read(workbook(H + ["mystery [kW]"], [["AHU-1", AC, 5, 2, 9]]), pack)
    assert any(e.column == "mystery" for e in r.errors)
    assert len(r.rows) == 1
    assert not any(i.name == "mystery" for row in r.rows for i in row.inputs)


def test_input_from_other_edition_is_unknown(pack):
    h = ["system tag", "system type", "max_airside_component_airflow [L/s]"]
    r = read(workbook(h, [["AHU-1", AC, 5]]), pack, "NCC2022")
    assert any(e.column == "max_airside_component_airflow" for e in r.errors)


def test_formula_refused(pack):
    r = read(workbook(H, [["AHU-1", AC, "=1+1", 2]]), pack)
    assert r.rows == [] and "formula" in r.errors[0].message.lower()


def test_hidden_sheet_ignored(pack):
    def extra(wb):
        ws = wb.create_sheet("Secret")
        ws.append(H)
        ws.append(["X", AC, 1, 1])
        ws.sheet_state = "hidden"
    r = read(workbook(H, [["AHU-1", AC, 5, 2]], extra), pack)
    assert [x.tag for x in r.rows] == ["AHU-1"]


def test_hidden_systems_sheet_is_an_error(pack):
    wb = openpyxl.Workbook()
    wb.active.title = "Systems"
    wb.active.append(H)
    wb.create_sheet("Other")
    wb["Systems"].sheet_state = "hidden"
    buf = io.BytesIO()
    wb.save(buf)
    assert read(buf.getvalue(), pack).errors


def test_row_limit(pack):
    rows = [[f"S{i}", AC, 1, 1] for i in range(MAX_ROWS + 1)]
    r = read(workbook(H, rows), pack)
    assert r.rows == [] and "2000" in r.errors[0].message


def test_duplicate_tag_and_missing_tag_and_bad_type(pack):
    r = read(workbook(H, [["A", AC, 1, 1], ["A", AC, 1, 1], [None, AC, 1, 1], ["B", "nonsense", 1, 1]]), pack)
    assert [x.tag for x in r.rows] == ["A"] and len(r.errors) == 3


def test_enum_bool_cells_and_project_facts_are_not_schedule_columns(pack):
    h = ["system tag", "system type", "supply_air_quantity_variable"]
    ok = read(workbook(h, [["A", AC, True]]), pack)
    assert ok.errors == []
    assert {i.name: i.value for i in ok.rows[0].inputs} == {"supply_air_quantity_variable": True}
    bad = read(workbook(h, [["A", AC, "maybe"]]), pack)
    assert bad.rows == [] and len(bad.errors) == 1
    facts = read(workbook(["system tag", "system type", "climate_zone [dimensionless]"], [["A", AC, 5]]), pack)
    assert any("climate_zone" in e.message for e in facts.errors)   # project facts come from the project


def test_only_xlsx_no_macros_or_external_links(pack):
    assert read(b"not a zip", pack).errors
    assert read(b"", pack).errors
    good = workbook(H, [["A", AC, 1, 1]])
    for name in ("xl/vbaProject.bin", "xl/externalLinks/externalLink1.xml"):
        src = zipfile.ZipFile(io.BytesIO(good))
        out = io.BytesIO()
        with zipfile.ZipFile(out, "w") as z:
            for n in src.namelist():
                z.writestr(n, src.read(n))
            z.writestr(name, b"x")
        r = read(out.getvalue(), pack)
        assert r.rows == [] and r.errors, name


def test_reads_a_path(pack, tmp_path):
    p = tmp_path / "s.xlsx"
    p.write_bytes(workbook(H, [["A", AC, 1, 1]]))
    assert len(read_schedule(p, pack, "NCC2022").rows) == 1
    (tmp_path / "s.xlsm").write_bytes(p.read_bytes())
    assert read_schedule(tmp_path / "s.xlsm", pack, "NCC2022").errors


def test_unknown_edition(pack):
    with pytest.raises(ValueError):
        build_template(pack, "NCC1999")
