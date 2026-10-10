"""Commissioning sheets and the re-import of site readings. Pure checks first, then the API on a signed revision. Needs the local Supabase."""
import io

import pytest
from fastapi.testclient import TestClient
from mep.api.server import create_pg_app
from mep.engine.loader import load_pack
from mep.review import commissioning as cx
from openpyxl import load_workbook
from pypdf import PdfReader

from tests.rls import test_pg_api as h
from tests.rls import test_pg_lineage as lin
from tests.rls import test_pg_review_api as rv
from tests.rls.conftest import DB_URL

TERMINALS = [{"tag": "T1", "system_tag": "SA-1", "airflow_ls": 200, "quantity": 2, "space_id": None}, {"tag": "=EVIL()", "system_tag": "SA-1", "airflow_ls": 100, "quantity": 1, "space_id": None},
             {"tag": "T9", "system_tag": "SA-2", "airflow_ls": None, "quantity": 1, "space_id": None}]


def filled(rows, readings, tweak=None):
    wb = load_workbook(io.BytesIO(cx.to_xlsx(rows, "rev-1")))
    ws = wb[cx.SHEET]
    for r in range(2, ws.max_row + 1):
        tag = ws.cell(r, 2).value
        tag = tag[1:] if isinstance(tag, str) and tag.startswith("'") else tag
        if tag in readings:
            ws.cell(r, 5).value = readings[tag]
    if tweak:
        tweak(ws)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def test_rows_one_per_unit_without_airflow_skipped_and_formulas_neutralised():
    rows = cx.rows_of(TERMINALS, {})
    assert [r["terminal"] for r in rows] == ["=EVIL()", "T1-1", "T1-2"] or [r["terminal"] for r in rows] == ["T1-1", "T1-2", "=EVIL()"]
    wb = load_workbook(io.BytesIO(cx.to_xlsx(rows, "rev-1")))
    cells = [str(c.value) for row in wb[cx.SHEET].iter_rows(min_row=2) for c in row if c.value is not None]
    assert not any(v.startswith(("=", "+", "@")) for v in cells) and any("EVIL" in v for v in cells)
    assert wb[cx.SHEET]["E2"].value is None and wb[cx.SHEET]["H2"].value is None                       # blank measured columns
    pdf = cx.to_pdf(rows, "rev-1")
    text = " ".join("\n".join(p.extract_text() or "" for p in PdfReader(io.BytesIO(pdf)).pages).split())
    assert cx.BANNER in text and "Measured L/s" in text and cx.to_pdf(rows, "rev-1") == pdf


def test_import_flags_by_the_given_tolerance_and_never_assumes_one():
    rows = cx.rows_of([TERMINALS[0]], {})
    sheet = filled(rows, {"T1-1": 209, "T1-2": 150})
    got = {r["terminal"]: r for r in cx.parse_import(sheet, rows, 5)}
    assert got["T1-1"]["flag"] == "WITHIN_TOLERANCE" and got["T1-1"]["variance_pct"] == 4.5
    assert got["T1-2"]["flag"] == "OUTSIDE_TOLERANCE" and got["T1-2"]["variance_pct"] == -25.0
    assert {r["flag"] for r in cx.parse_import(sheet, rows, 30)} == {"WITHIN_TOLERANCE"}
    for bad in (-1, 51, float("nan")):
        with pytest.raises(cx.ImportRefused):
            cx.parse_import(sheet, rows, bad)


def test_import_catches_a_changed_design_bad_values_and_a_foreign_file():
    rows = cx.rows_of([TERMINALS[0]], {})

    def change_design(ws):
        ws["D2"].value = 999

    flags = {r["terminal"]: r["flag"] for r in cx.parse_import(filled(rows, {"T1-1": 200, "T1-2": 200}, change_design), rows, 5)}
    assert flags == {"T1-1": "DESIGN_CHANGED", "T1-2": "WITHIN_TOLERANCE"}
    for value, flag in (("abc", "INVALID"), (-5, "INVALID"), (True, "INVALID"), (None, "NOT_MEASURED")):
        got = {r["terminal"]: r["flag"] for r in cx.parse_import(filled(rows, {"T1-1": value, "T1-2": 200}), rows, 5)}
        assert got["T1-1"] == flag
    extra = {r["terminal"]: r["flag"] for r in cx.parse_import(filled(rows, {"T1-1": 200}, lambda ws: ws.append(["SA-1", "T7", "", 50, 50, "", "", ""])), rows, 5)}
    assert extra["T7"] == "NOT_IN_DESIGN" and extra["T1-2"] == "NOT_MEASURED"
    dup = [r["flag"] for r in cx.parse_import(filled(rows, {"T1-1": 200}, lambda ws: ws.append(["SA-1", "T1-1", "", 200, 200, "", "", ""])), rows, 5)]
    assert "DUPLICATE" in dup
    assert "MISSING_FROM_SHEET" in {r["flag"] for r in cx.parse_import(filled(rows, {}, lambda ws: ws.delete_rows(3)), rows, 5)}
    for junk in (b"not a workbook", b""):
        with pytest.raises(cx.ImportRefused):
            cx.parse_import(junk, rows, 5)

    def renamed(ws):
        ws["D1"].value = "Design"

    with pytest.raises(cx.ImportRefused):
        cx.parse_import(filled(rows, {}, renamed), rows, 5)
    with pytest.raises(cx.ImportRefused):
        cx.parse_import(b"x" * (cx.MAX_IMPORT_BYTES + 1), rows, 5)


@pytest.fixture(scope="module")
def client():
    return TestClient(create_pg_app(DB_URL, h.SECRET, load_pack(lin.ROOT / "rules"), supabase_url=lin.SUPABASE_URL, anon_key=lin.ANON))


def test_sheets_need_a_signed_revision_and_readings_are_recorded_not_resulted(admin, client):
    pack = load_pack(lin.ROOT / "rules")
    f = rv.frozen(admin, client, pack)
    rev, de = f["revision"], rv.auth(f, "designer")
    assert client.get(f"/revisions/{rev}/commissioning.xlsx", headers=de).json()["detail"]["code"] == "not_signed"
    rv.review_all(client, f)
    assert client.post(f"/revisions/{rev}/sign/gate2", headers=rv.auth(f, "checker"), json={}).status_code == 200
    rv.acknowledge_all(client, f)
    assert client.post(f"/revisions/{rev}/sign/gate3", headers=rv.auth(f, "approver"), json={"registration": "RPEQ 12345"}).status_code == 200
    assert client.get(f"/revisions/{rev}/commissioning.xlsx", headers=de).json()["detail"]["code"] == "no_terminals"
    admin.execute("set session_replication_role = replica")
    admin.execute("insert into duct_run (firm_id, revision_id, kind, tag, system_tag, quantity, airflow_ls, created_by) values (%s, %s, 'terminal', 'T1', 'SA-1', 2, 200, %s)",
                  (f["firm"], rev, f["designer"]))
    admin.execute("set session_replication_role = origin")
    xlsx = client.get(f"/revisions/{rev}/commissioning.xlsx", headers=de)
    assert xlsx.status_code == 200
    assert client.get(f"/revisions/{rev}/commissioning.pdf", headers=de).content[:5] == b"%PDF-"
    sheet = load_workbook(io.BytesIO(xlsx.content))
    ws = sheet[cx.SHEET]
    ws["E2"].value, ws["E3"].value = 205, 120
    buf = io.BytesIO()
    sheet.save(buf)
    url = f"/revisions/{rev}/commissioning/import"
    assert client.post(url, headers=rv.auth(f, "checker"), data={"tolerance_pct": "10"}, files={"file": ("s.xlsx", buf.getvalue())}).status_code == 403
    assert client.post(url, headers=de, data={"tolerance_pct": "90"}, files={"file": ("s.xlsx", buf.getvalue())}).status_code == 422
    assert client.post(url, headers=de, data={"tolerance_pct": "10"}, files={"file": ("s.xlsx", b"junk")}).status_code == 422
    ok = client.post(url, headers=de, data={"tolerance_pct": "10"}, files={"file": ("s.xlsx", buf.getvalue())})
    assert ok.status_code == 200 and ok.json()["counts"] == {"WITHIN_TOLERANCE": 1, "OUTSIDE_TOLERANCE": 1}
    stored = client.get(f"/revisions/{rev}/commissioning/readings", headers=de).json()["batches"]
    assert len(stored) == 1 and len(stored[0]["readings"]) == 2
    assert admin.execute("select count(*) from ledger_event where revision_id = %s and kind = 'commissioning_imported'", (rev,)).fetchone()[0] == 1
    assert admin.execute("select count(*) from rule_result where revision_id = %s and rule_id like 'COMM%%'", (rev,)).fetchone()[0] == 0
    other = h.seed(admin)
    assert client.get(f"/revisions/{rev}/commissioning/readings", headers=h.auth(other["designer"])).status_code == 404
