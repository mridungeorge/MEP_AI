"""Services schedule, ceiling-void check and quantities. Needs the local Supabase."""
import csv
import io

import pytest
from fastapi.testclient import TestClient
from mep.api.server import create_pg_app
from mep.engine.loader import load_pack
from openpyxl import load_workbook

from tests.rls import test_pg_api as h
from tests.rls import test_pg_lineage as lin
from tests.rls.conftest import DB_URL

PACK = load_pack(lin.ROOT / "rules")


@pytest.fixture(scope="module")
def client():
    return TestClient(create_pg_app(DB_URL, h.SECRET, PACK, supabase_url=lin.SUPABASE_URL, anon_key=lin.ANON))


def setup(admin, client):
    f = h.seed(admin)
    d = h.auth(f["designer"])
    for name, void in (("Office", 400), ("Store", None)):
        body = {"name": name, "area_m2": 20, "storey": "Level 1", "use": "office"}
        if void is not None:
            body["ceiling_void_mm"] = void
        assert client.post(f"{h.base(f)}/gate1/spaces", json=body, headers=d).status_code == 200
    spaces = {s["name"]: s["id"] for s in client.get(f"{h.base(f)}/gate1", headers=d).json()["spaces"]}
    return f, d, spaces


def duct(tag, space, **kw):
    return {"kind": "duct", "tag": tag, "space_id": space, "shape": "rect", "width": 600, "depth": 300, "length": 4, "insulation": 25, "quantity": 1, **kw}


def test_the_schedule_is_a_designers_to_edit_and_converts_units(admin, client):
    f, d, sp = setup(admin, client)
    url = f"{h.base(f)}/services"
    r = client.post(url, json=duct("D1", sp["Office"], section_unit="in", width=24, depth=12, length=10, length_unit="ft"), headers=d)
    assert r.status_code == 200
    item = client.get(url, headers=d).json()["items"][0]
    assert item["width_mm"] == 609.6 and item["depth_mm"] == 304.8 and item["length_m"] == 3.048
    assert client.post(url, json=duct("D2", sp["Office"]), headers=h.auth(f["checker"])).status_code == 403
    assert client.post(url, json={"kind": "duct", "tag": "bad"}, headers=d).status_code == 422                       # no section or length
    assert client.post(url, json=duct("D3", "00000000-0000-0000-0000-000000000000"), headers=d).status_code == 422    # not a space of this revision
    assert client.post(url, json=duct("D4", sp["Office"], width=-5), headers=d).status_code == 422
    assert client.post(url, json={**duct("D5", None), "extra": 1}, headers=d).status_code == 422
    assert client.put(f"{url}/{r.json()['id']}", json=duct("D1", sp["Office"], depth=250), headers=d).status_code == 200
    assert client.delete(f"{url}/{r.json()['id']}", headers=d).status_code == 200
    assert client.get(url, headers=h.auth(h.seed(admin)["designer"])).status_code == 404                              # another firm cannot see the revision


def test_the_ceiling_void_check_adds_insulation_on_both_faces_and_the_clearance(admin, client):
    f, d, sp = setup(admin, client)
    url = f"{h.base(f)}/services"
    client.post(url, json=duct("D1", sp["Office"], depth=300, insulation=25), headers=d)           # 300 + 2x25 + 50 = 400 -> exactly fits 400
    client.post(url, json=duct("D2", sp["Store"]), headers=d)                                          # the Store declares no void
    rows = {r["space"]: r for r in client.get(f"{h.base(f)}/ceiling-void", headers=d).json()["spaces"]}
    assert rows["Office"]["required_mm"] == 400 and rows["Office"]["status"] == "CLEAR" and rows["Office"]["margin_mm"] == 0
    assert rows["Store"]["status"] == "NO DATA"
    client.post(url, json=duct("D3", sp["Office"], depth=310), headers=d)
    rows = {r["space"]: r for r in client.get(f"{h.base(f)}/ceiling-void", headers=d).json()["spaces"]}
    assert rows["Office"]["status"] == "CLASH" and rows["Office"]["deepest_duct"] == "D3" and rows["Office"]["margin_mm"] == -10
    assert "PASS" not in str(client.get(f"{h.base(f)}/ceiling-void", headers=d).json())


def test_the_clearance_is_a_firm_setting(admin, client):
    f, d, sp = setup(admin, client)
    client.post(f"{h.base(f)}/services", json=duct("D1", sp["Office"], depth=300, insulation=25), headers=d)
    admin.execute("update app_user set is_admin = true where id = %s", (f["designer"],))
    a = d
    firm = client.get("/admin/overview", headers=a).json()["firm"]
    assert firm["void_clearance_mm"] == 50
    body = {k: firm[k] for k in ("name", "signer_mode", "sample_size", "near_miss_default")} | {"void_clearance_mm": 80}
    assert client.put("/admin/firm", json=body, headers=a).status_code == 200
    assert client.put("/admin/firm", json=body | {"void_clearance_mm": 5000}, headers=a).status_code == 422
    assert client.get(f"{h.base(f)}/ceiling-void", headers=d).json()["clearance_mm"] == 80
    row = next(r for r in client.get(f"{h.base(f)}/ceiling-void", headers=d).json()["spaces"] if r["space"] == "Office")
    assert row["status"] == "CLASH"


def test_quantities_by_size_fittings_insulation_and_terminals(admin, client):
    f, d, sp = setup(admin, client)
    url = f"{h.base(f)}/services"
    client.post(url, json=duct("D1", sp["Office"], width=600, depth=300, length=5, insulation=25, quantity=2), headers=d)    # 2*(0.9*2)*5 = 18 m2
    client.post(url, json=duct("D2", sp["Office"], shape="round", diameter=250, width=None, depth=None, length=2, insulation=0), headers=d)
    client.post(url, json={"kind": "fitting", "tag": "F1", "fitting_type": "90 degree bend", "quantity": 3}, headers=d)
    client.post(url, json={"kind": "terminal", "tag": "T1", "space_id": sp["Office"], "quantity": 4}, headers=d)
    client.post(url, json={"kind": "terminal", "tag": "T2", "quantity": 1}, headers=d)
    q = client.get(f"{h.base(f)}/quantities", headers=d).json()
    rect = next(x for x in q["duct_by_size"] if x["shape"] == "rect")
    assert rect["count"] == 2 and rect["length_m"] == 10 and rect["surface_area_m2"] == 18
    assert next(x for x in q["duct_by_size"] if x["shape"] == "round")["surface_area_m2"] == pytest.approx(1.571, abs=0.001)
    assert q["insulation"] == [{"thickness_mm": 25.0, "area_m2": 18.0}]
    assert q["fittings"] == [{"type": "90 degree bend", "quantity": 3}]
    assert {t["space"]: t["quantity"] for t in q["terminals_per_space"]} == {"Office": 4, "(no space)": 1}
    rows = list(csv.reader(io.StringIO(client.get(f"{h.base(f)}/quantities.csv", headers=d).text)))
    assert ["duct_by_size"] in rows and any("90 degree bend" in r for r in rows)
    wb = load_workbook(io.BytesIO(client.get(f"{h.base(f)}/quantities.xlsx", headers=d).content))
    assert {"duct_by_size", "fittings", "insulation", "terminals_per_space"} <= set(wb.sheetnames)


def test_a_frozen_revision_takes_no_change(admin, client):
    f, d, sp = setup(admin, client)
    admin.execute("update revision set frozen_at = now(), status = 'frozen' where id = %s", (f["revision"],))
    assert client.post(f"{h.base(f)}/services", json=duct("D1", sp["Office"]), headers=d).status_code == 409
