"""Performance Solution pathway: the flag, the pathway choice, engineer evidence (never a rule result) and the starting data package. Needs the local Supabase."""
import io
import json

import psycopg
import pytest
from fastapi.testclient import TestClient
from mep.api import performance
from mep.api.server import create_pg_app
from mep.engine.loader import load_pack
from openpyxl import load_workbook

from tests.rls import test_pg_api as h
from tests.rls import test_pg_lineage as lin
from tests.rls.conftest import DB_URL

PACK = load_pack(lin.ROOT / "rules")
SUBJECT, RULE = "ahu-1", "NCC2025-J6D3-time-switch-ac"


@pytest.fixture(scope="module")
def client():
    return TestClient(create_pg_app(DB_URL, h.SECRET, PACK, supabase_url=lin.SUPABASE_URL, anon_key=lin.ANON))


def run_revision(admin, client):
    f = h.seed(admin)
    h.populate(client, PACK, f)
    assert h.confirm(client, f, h.rows_to_confirm(client, f)).status_code == 200
    assert client.post(f"/revisions/{f['revision']}/run-rules", headers=h.auth(f["designer"])).status_code == 200
    return f


def base(f):
    return f"/revisions/{f['revision']}/performance"


def test_the_flag_follows_the_engine_and_the_pathway_is_the_designers_choice(admin, client, monkeypatch):
    f = run_revision(admin, client)
    d = h.auth(f["designer"])
    row = next(r for r in client.get(base(f), headers=d).json()["results"] if r["rule_id"] == RULE)
    assert row["pathway"] == "DTS" and row["performance_solution_likely"] is False          # the engine does find an option for this one
    monkeypatch.setattr(performance, "_likely", lambda *a, **k: True)                            # a FAIL with no passing option
    row = next(r for r in client.get(base(f), headers=d).json()["results"] if r["rule_id"] == RULE)
    assert row["flag"] == "Performance Solution pathway likely"
    url = f"{base(f)}/{SUBJECT}/{RULE}/pathway"
    assert client.put(url, headers=h.auth(f["checker"]), json={"pathway": "PERFORMANCE_SOLUTION"}).status_code == 403
    assert client.put(url, headers=d, json={"pathway": "MAYBE"}).status_code == 422
    assert client.put(f"{base(f)}/{SUBJECT}/NCC2025-NOPE/pathway", headers=d, json={"pathway": "PERFORMANCE_SOLUTION"}).status_code == 422
    assert client.put(url, headers=d, json={"pathway": "PERFORMANCE_SOLUTION", "note": "needs a thermal model"}).status_code == 200
    row = next(r for r in client.get(base(f), headers=d).json()["results"] if r["rule_id"] == RULE)
    assert row["pathway"] == "PERFORMANCE_SOLUTION" and row["note"] == "needs a thermal model"
    assert admin.execute("select count(*) from ledger_event where revision_id = %s and kind = 'result_pathway_set'", (f["revision"],)).fetchone()[0] == 1


def test_evidence_is_recorded_as_evidence_and_never_changes_a_result(admin, client):
    f = run_revision(admin, client)
    d = h.auth(f["designer"])
    before = admin.execute("select count(*), string_agg(result::text, ',' order by rule_id) from rule_result where revision_id = %s", (f["revision"],)).fetchone()
    url = f"{base(f)}/{SUBJECT}/{RULE}/evidence"
    metrics = json.dumps([{"name": "peak operative temperature", "value": 26.4, "unit": "degC"}])
    ok = client.post(url, headers=d, data={"title": "Thermal model run 3", "tool": "IES VE", "description": "annual simulation", "metrics": metrics},
                     files={"file": ("model-report.pdf", b"%PDF-1.4 fake report")})
    assert ok.status_code == 200 and ok.json()["provenance"] == "engineer_supplied"
    eid = ok.json()["evidence_id"]
    for bad in ('[{"name": "x", "value": "hot", "unit": "degC"}]', '[{"name": "x", "value": 1, "unit": "nonsense_unit_xyz"}]', "not json",
                '[{"name": "x", "value": 1}]'):
        assert client.post(url, headers=d, data={"title": "bad one", "metrics": bad}).status_code == 422
    assert client.post(url, headers=h.auth(f["checker"]), data={"title": "not mine"}).status_code == 403
    assert client.post(f"{base(f)}/{SUBJECT}/NCC2025-NOPE/evidence", headers=d, data={"title": "no such result"}).status_code == 422
    after = admin.execute("select count(*), string_agg(result::text, ',' order by rule_id) from rule_result where revision_id = %s", (f["revision"],)).fetchone()
    assert after == before                                                                        # no result was created or changed
    assert admin.execute("select count(*) from perf_evidence where revision_id = %s", (f["revision"],)).fetchone()[0] == 1
    got = client.get(f"{base(f)}/evidence/{eid}/file", headers=d)
    assert got.status_code == 200 and got.content == b"%PDF-1.4 fake report"
    other = h.seed(admin)
    assert client.get(f"{base(f)}/evidence/{eid}/file", headers=h.auth(other["designer"])).status_code == 404
    with pytest.raises(psycopg.errors.Error):                                                       # append-only
        admin.execute("update perf_evidence set title = 'edited' where id = %s", (eid,))
    ev = client.get(base(f), headers=d).json()["results"]
    assert any(e["id"] == eid for r in ev for e in r["evidence"])
    with psycopg.connect(DB_URL, autocommit=True) as c:
        c.execute("set role authenticated")
        c.execute("select set_config('request.jwt.claims', %s, false)", (json.dumps({"sub": str(f["designer"])}),))
        with pytest.raises(psycopg.errors.InsufficientPrivilege):
            c.execute("select file_content from perf_evidence")
    admin.execute("update revision set frozen_at = now(), status = 'frozen' where id = %s", (f["revision"],))
    assert client.post(url, headers=d, data={"title": "late evidence"}).status_code == 409


def test_the_starting_package_holds_the_failed_benchmarks_and_is_safe_to_open(admin, client):
    f = run_revision(admin, client)
    d = h.auth(f["designer"])
    client.post(f"{base(f)}/{SUBJECT}/{RULE}/evidence", headers=d, data={"title": "=HYPERLINK(\"http://evil.example\")", "tool": "+cmd", "metrics": "[]"})
    pkg = client.get(f"{base(f)}/package.json", headers=d).json()
    assert pkg["banner"].startswith("STARTING DATA PACKAGE") and "DRAFT RULES" in pkg["banner"]
    failed = {(b["subject_id"], b["rule_id"]): b for b in pkg["failed_dts_benchmarks"]}
    assert (SUBJECT, RULE) in failed and failed[(SUBJECT, RULE)]["dts_check"] and failed[(SUBJECT, RULE)]["citation"]["clause"]
    assert {"project", "spaces", "systems", "evidence"} <= set(pkg) and pkg["systems"][0]["tag"] == SUBJECT and pkg["systems"][0]["inputs"]
    assert all(b["rule_status"] == "draft" for b in pkg["failed_dts_benchmarks"])
    x = client.get(f"{base(f)}/package.xlsx", headers=d)
    assert x.status_code == 200
    wb = load_workbook(io.BytesIO(x.content))
    assert {"README", "Project", "Spaces", "Systems", "Failed DTS benchmarks", "Evidence"} <= set(wb.sheetnames)
    cells = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    assert not any(v.startswith(("=", "+", "@")) for v in cells)                                     # nothing can run when the file is opened
    assert any("HYPERLINK" in v for v in cells)                                                        # but the text is kept
    assert client.get(f"{base(f)}/package.json", headers=h.auth(h.seed(admin)["designer"])).status_code == 404
