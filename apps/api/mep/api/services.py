"""The mechanical services schedule: proposed ducts, fittings and terminals; the ceiling-void check; quantities.

None of this is compliance. It is arithmetic on figures the designer entered (every length goes through pint to millimetres or metres) so that a clash with the
ceiling void, or a take-off, is visible early. Results say CLASH / CLEAR / NO DATA, never PASS or FAIL.
"""
import csv
import io
import math
from collections.abc import Iterator
from contextlib import contextmanager
from typing import Annotated, Any, Literal
from uuid import UUID

import psycopg
from fastapi import APIRouter, Depends, HTTPException
from fastapi.responses import Response
from psycopg.rows import dict_row
from pydantic import BaseModel, ConfigDict, Field, model_validator

from mep.api.revisions import DESIGNER_ROLES, Repo, _err
from mep.api.revisions import User as RevUser
from mep.engine import units

router = APIRouter()
LENGTH_UNITS = Literal["mm", "m", "in", "ft"]


def get_dsn() -> str:
    raise HTTPException(status_code=503, detail="the services schedule is not configured")


Dsn = Annotated[str, Depends(get_dsn)]


@contextmanager
def _as_user(dsn: str, user: Any) -> Iterator[psycopg.Connection[dict[str, Any]]]:
    with psycopg.connect(dsn, autocommit=False, row_factory=dict_row) as conn:
        conn.execute("set local role authenticated")
        conn.execute("select set_config('request.jwt.claims', %s, true)", (user.claims_json(),))
        yield conn


def to_mm(value: float | None, unit: str) -> float | None:
    """A length in `unit` as millimetres, through the project's pint registry."""
    if value is None:
        return None
    if not math.isfinite(value):
        raise _err(422, "bad_number", "numbers must be finite")
    return round(float((value * units.unit_of(unit)).to("mm").magnitude), 3)


class RunBody(BaseModel):
    model_config = ConfigDict(extra="forbid")

    kind: Literal["duct", "fitting", "terminal"]
    tag: str = Field(min_length=1, max_length=60)
    system_tag: str | None = Field(default=None, max_length=60)
    space_id: UUID | None = None
    shape: Literal["rect", "round"] | None = None
    width: float | None = Field(default=None, gt=0)
    depth: float | None = Field(default=None, gt=0)
    diameter: float | None = Field(default=None, gt=0)
    section_unit: LENGTH_UNITS = "mm"
    length: float | None = Field(default=None, gt=0)
    length_unit: LENGTH_UNITS = "m"
    insulation: float = Field(default=0, ge=0)
    insulation_unit: LENGTH_UNITS = "mm"
    fitting_type: str | None = Field(default=None, max_length=60)
    quantity: int = Field(default=1, ge=1, le=10000)
    airflow_ls: float | None = Field(default=None, ge=0, le=1_000_000)
    start_mm: tuple[float, float, float] | None = None
    end_mm: tuple[float, float, float] | None = None

    @model_validator(mode="after")
    def _complete(self) -> "RunBody":
        if self.kind == "duct":
            ok = self.length is not None and ((self.shape == "rect" and self.width and self.depth) or (self.shape == "round" and self.diameter))
            if not ok:
                raise ValueError("a duct needs a shape, a length and its section (width and depth, or a diameter)")
        if self.kind == "fitting" and not self.fitting_type:
            raise ValueError("a fitting needs its type")
        if (self.start_mm is None) != (self.end_mm is None):
            raise ValueError("give both ends of the run or neither")
        return self


def _columns(b: RunBody) -> dict[str, Any]:
    start, end = b.start_mm or (None, None, None), b.end_mm or (None, None, None)
    return {"kind": b.kind, "tag": b.tag.strip(), "system_tag": b.system_tag, "space_id": b.space_id, "shape": b.shape if b.kind == "duct" else None,
            "width_mm": to_mm(b.width, b.section_unit) if b.kind == "duct" and b.shape == "rect" else None,
            "depth_mm": to_mm(b.depth, b.section_unit) if b.kind == "duct" and b.shape == "rect" else None,
            "diameter_mm": to_mm(b.diameter, b.section_unit) if b.kind == "duct" and b.shape == "round" else None,
            "length_m": round(float(to_mm(b.length, b.length_unit) or 0) / 1000, 4) if b.kind == "duct" and b.length else None,
            "insulation_mm": to_mm(b.insulation, b.insulation_unit) or 0, "fitting_type": b.fitting_type if b.kind == "fitting" else None, "quantity": b.quantity,
            "airflow_ls": b.airflow_ls, "x0": start[0], "y0": start[1], "z0": start[2], "x1": end[0], "y1": end[1], "z1": end[2]}


def _open_revision(repo: Any, user: Any, revision_id: UUID) -> None:
    info = repo.revision_info(revision_id, user.firm_id)
    if info is None:
        raise _err(404, "not_found", "revision not found")
    if user.role not in DESIGNER_ROLES:
        raise _err(403, "forbidden", "only a designer edits the services schedule")
    if info["frozen"]:
        raise _err(409, "revision_frozen", "the revision is frozen")


def _check_space(dsn: str, user: Any, revision_id: UUID, space_id: UUID | None) -> None:
    if space_id is None:
        return
    with _as_user(dsn, user) as conn:
        if conn.execute("select 1 from space where id = %s and revision_id = %s and firm_id = %s", (space_id, revision_id, user.firm_id)).fetchone() is None:
            raise _err(422, "bad_space", "that space is not in this revision")


def _service_write(dsn: str, sql: str, args: tuple[Any, ...]) -> dict[str, Any] | None:
    try:
        with psycopg.connect(dsn, autocommit=True, row_factory=dict_row) as conn:
            return conn.execute(sql, args).fetchone()
    except psycopg.errors.RaiseException as exc:
        raise _err(409 if "frozen" in str(exc) else 422, "refused", str(exc).splitlines()[0]) from None
    except psycopg.errors.CheckViolation as exc:
        raise _err(422, "invalid", str(exc).splitlines()[0]) from None


@router.post("/revisions/{revision_id}/services")
def add_run(revision_id: UUID, body: RunBody, user: RevUser, repo: Repo, dsn: Dsn) -> dict[str, Any]:
    _open_revision(repo, user, revision_id)
    _check_space(dsn, user, revision_id, body.space_id)
    c = _columns(body)
    names = list(c)
    row = _service_write(dsn, f"insert into duct_run (firm_id, revision_id, created_by, {', '.join(names)}) values (%s, %s, %s, {', '.join(['%s'] * len(names))}) returning id",     # noqa: S608
                         (user.firm_id, revision_id, user.user_id, *c.values()))
    assert row is not None
    return {"id": str(row["id"])}


@router.put("/revisions/{revision_id}/services/{run_id}")
def update_run(revision_id: UUID, run_id: UUID, body: RunBody, user: RevUser, repo: Repo, dsn: Dsn) -> dict[str, Any]:
    _open_revision(repo, user, revision_id)
    _check_space(dsn, user, revision_id, body.space_id)
    c = _columns(body)
    row = _service_write(dsn, f"update duct_run set {', '.join(f'{k} = %s' for k in c)} where id = %s and revision_id = %s and firm_id = %s returning id",       # noqa: S608
                         (*c.values(), run_id, revision_id, user.firm_id))
    if row is None:
        raise _err(404, "not_found", "no such item")
    return {"id": str(row["id"])}


@router.delete("/revisions/{revision_id}/services/{run_id}")
def delete_run(revision_id: UUID, run_id: UUID, user: RevUser, repo: Repo, dsn: Dsn) -> dict[str, Any]:
    _open_revision(repo, user, revision_id)
    row = _service_write(dsn, "delete from duct_run where id = %s and revision_id = %s and firm_id = %s returning id", (run_id, revision_id, user.firm_id))
    if row is None:
        raise _err(404, "not_found", "no such item")
    return {"deleted": True}


def runs_of(dsn: str, user: Any, revision_id: UUID) -> list[dict[str, Any]]:
    with _as_user(dsn, user) as conn:
        rows = conn.execute("select * from duct_run where revision_id = %s and firm_id = %s order by kind, tag, created_at", (revision_id, user.firm_id)).fetchall()
    out = []
    for r in rows:
        out.append({k: (str(v) if k in ("id", "firm_id", "revision_id", "space_id", "created_by") and v is not None else (float(v) if hasattr(v, "as_tuple") else v))
                    for k, v in r.items() if k not in ("firm_id", "created_by", "created_at")})
    return out


@router.get("/revisions/{revision_id}/services")
def list_runs(revision_id: UUID, user: RevUser, repo: Repo, dsn: Dsn) -> dict[str, Any]:
    if repo.revision_info(revision_id, user.firm_id) is None:
        raise _err(404, "not_found", "revision not found")
    return {"items": runs_of(dsn, user, revision_id)}


# ---------------------------------------------------------------------------------------------------------------------------- ceiling void
def section_depth(r: dict[str, Any]) -> float:
    return float(r["depth_mm"]) if r["shape"] == "rect" else float(r["diameter_mm"])


@router.get("/revisions/{revision_id}/ceiling-void")
def ceiling_void(revision_id: UUID, user: RevUser, repo: Repo, dsn: Dsn) -> dict[str, Any]:
    """Per space: the deepest proposed duct, its insulation on both faces and the firm's clearance, against the ceiling void the space declares."""
    if repo.revision_info(revision_id, user.firm_id) is None:
        raise _err(404, "not_found", "revision not found")
    with _as_user(dsn, user) as conn:
        clearance = float(conn.execute("select void_clearance_mm from firm where id = %s", (user.firm_id,)).fetchone()["void_clearance_mm"])
        spaces = conn.execute("select id, name, ceiling_void_mm_value as void from space where revision_id = %s and firm_id = %s order by name", (revision_id, user.firm_id)).fetchall()
    ducts = [r for r in runs_of(dsn, user, revision_id) if r["kind"] == "duct" and r["space_id"]]
    rows = []
    for s in spaces:
        mine = [d for d in ducts if d["space_id"] == str(s["id"])]
        void = None if s["void"] is None else float(s["void"])
        if not mine:
            rows.append({"space_id": str(s["id"]), "space": s["name"], "ceiling_void_mm": void, "status": "NO DATA", "reason": "no duct is assigned to this space"})
            continue
        worst = max(mine, key=lambda d: section_depth(d) + 2 * float(d["insulation_mm"]))
        need = section_depth(worst) + 2 * float(worst["insulation_mm"]) + clearance
        if void is None:
            rows.append({"space_id": str(s["id"]), "space": s["name"], "ceiling_void_mm": None, "status": "NO DATA", "reason": "the space has no ceiling void figure", "deepest_duct": worst["tag"]})
            continue
        rows.append({"space_id": str(s["id"]), "space": s["name"], "ceiling_void_mm": void, "deepest_duct": worst["tag"], "duct_depth_mm": section_depth(worst),
                     "insulation_each_face_mm": float(worst["insulation_mm"]), "clearance_mm": clearance, "required_mm": need, "margin_mm": round(void - need, 3),
                     "status": "CLASH" if need > void else "CLEAR"})
    return {"note": "Warnings only. The deepest single duct is checked; crossings and stacked services are not modelled. A CLASH is a prompt to check, not a finding.",
            "clearance_mm": clearance, "spaces": rows}


# ---------------------------------------------------------------------------------------------------------------------------- quantities
def quantities_of(dsn: str, user: Any, revision_id: UUID) -> dict[str, Any]:
    items = runs_of(dsn, user, revision_id)
    with _as_user(dsn, user) as conn:
        names = {str(s["id"]): s["name"] for s in conn.execute("select id, name from space where revision_id = %s and firm_id = %s", (revision_id, user.firm_id)).fetchall()}
    by_size: dict[tuple[str, str], dict[str, float]] = {}
    insulation: dict[float, float] = {}
    fittings: dict[str, int] = {}
    terminals: dict[str, int] = {}
    for r in items:
        if r["kind"] == "duct":
            if r["shape"] == "rect":
                size, perimeter_m = f"{r['width_mm']:g} x {r['depth_mm']:g}", 2 * (float(r["width_mm"]) + float(r["depth_mm"])) / 1000
            else:
                size, perimeter_m = f"dia {r['diameter_mm']:g}", math.pi * float(r["diameter_mm"]) / 1000
            area = perimeter_m * float(r["length_m"]) * int(r["quantity"])
            slot = by_size.setdefault((r["shape"], size), {"count": 0, "length_m": 0.0, "area_m2": 0.0})
            slot["count"] += int(r["quantity"])
            slot["length_m"] += float(r["length_m"]) * int(r["quantity"])
            slot["area_m2"] += area
            if float(r["insulation_mm"]) > 0:
                insulation[float(r["insulation_mm"])] = insulation.get(float(r["insulation_mm"]), 0.0) + area
        elif r["kind"] == "fitting":
            fittings[r["fitting_type"]] = fittings.get(r["fitting_type"], 0) + int(r["quantity"])
        else:
            key = names.get(r["space_id"] or "", "(no space)")
            terminals[key] = terminals.get(key, 0) + int(r["quantity"])
    return {"duct_by_size": [{"shape": k[0], "size_mm": k[1], "count": int(v["count"]), "length_m": round(v["length_m"], 3), "surface_area_m2": round(v["area_m2"], 3)}
                             for k, v in sorted(by_size.items())],
            "fittings": [{"type": k, "quantity": v} for k, v in sorted(fittings.items())],
            "insulation": [{"thickness_mm": k, "area_m2": round(v, 3)} for k, v in sorted(insulation.items())],
            "terminals_per_space": [{"space": k, "quantity": v} for k, v in sorted(terminals.items())],
            "basis": "Duct surface area is the bare duct's perimeter times its length; insulation area is that surface for the ducts that carry insulation. Quantities follow the entered schedule only."}


@router.get("/revisions/{revision_id}/quantities")
def quantities(revision_id: UUID, user: RevUser, repo: Repo, dsn: Dsn) -> dict[str, Any]:
    if repo.revision_info(revision_id, user.firm_id) is None:
        raise _err(404, "not_found", "revision not found")
    return quantities_of(dsn, user, revision_id)


def _safe(v: Any) -> Any:
    return "'" + v if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r") else v


@router.get("/revisions/{revision_id}/quantities.csv")
def quantities_csv(revision_id: UUID, user: RevUser, repo: Repo, dsn: Dsn) -> Response:
    if repo.revision_info(revision_id, user.firm_id) is None:
        raise _err(404, "not_found", "revision not found")
    q = quantities_of(dsn, user, revision_id)
    out = io.StringIO()
    w = csv.writer(out)
    for section in ("duct_by_size", "fittings", "insulation", "terminals_per_space"):
        w.writerow([section])
        rows = q[section]
        if rows:
            w.writerow(list(rows[0]))
            for r in rows:
                w.writerow([_safe(v) for v in r.values()])
        w.writerow([])
    w.writerow([q["basis"]])
    return Response(content=out.getvalue(), media_type="text/csv", headers={"Content-Disposition": 'attachment; filename="quantities.csv"', "Cache-Control": "no-store"})


@router.get("/revisions/{revision_id}/quantities.xlsx")
def quantities_xlsx(revision_id: UUID, user: RevUser, repo: Repo, dsn: Dsn) -> Response:
    from openpyxl import Workbook
    if repo.revision_info(revision_id, user.firm_id) is None:
        raise _err(404, "not_found", "revision not found")
    q = quantities_of(dsn, user, revision_id)
    wb = Workbook()
    wb.active.title = "Basis"
    wb.active.append([q["basis"]])
    for section in ("duct_by_size", "fittings", "insulation", "terminals_per_space"):
        ws = wb.create_sheet(section[:31])
        rows = q[section]
        if rows:
            ws.append(list(rows[0]))
            for r in rows:
                ws.append([_safe(v) for v in r.values()])
    buf = io.BytesIO()
    wb.save(buf)
    return Response(content=buf.getvalue(), media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    headers={"Content-Disposition": 'attachment; filename="quantities.xlsx"', "Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"})
