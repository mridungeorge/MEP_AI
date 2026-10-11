"""Commissioning sheets: one row per terminal with its DESIGN airflow from the signed revision and blank columns for what a technician measures on site, and the
re-import of those measurements with a tolerance flag.

A measurement is a site reading, not a compliance result and not an input to any rule. The tolerance is supplied by the person importing (it is a project or
contract matter, never assumed here). Spreadsheet cells we write that start like a formula are stored as text; the import reads values only.
"""
import io
import math
import re
import unicodedata
import zipfile
from typing import Any
from xml.sax.saxutils import escape

from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

SHEET = "Terminals"
HEADER = ["System", "Terminal", "Space", "Design airflow (L/s)", "Measured airflow (L/s)", "Measured on", "Measured by", "Comment"]
COL_MEASURED = 4
MAX_IMPORT_BYTES = 5 * 1024 * 1024
MAX_ROWS = 5000
MAX_UNPACKED = 40 * 1024 * 1024
CONTROL = re.compile("[\\x00-\\x08\\x0b\\x0c\\x0e-\\x1f\\x7f]")
BANNER = "COMMISSIONING SHEET: site readings are records, not compliance results"


class ImportRefused(Exception):
    pass


def clean(v: Any) -> str:
    return CONTROL.sub("", str(v if v is not None else ""))


def norm(v: Any) -> str:
    """One spelling of a tag or system on both sides of the round trip: control characters out, NFKC, trimmed."""
    return unicodedata.normalize("NFKC", clean(v)).strip()


def put_text(cell: Any, v: Any) -> None:
    """Store text so a spreadsheet can never run it: a plain string cell with the quote prefix style, not a changed value."""
    cell.value = clean(v) if isinstance(v, str) else v
    if isinstance(cell.value, str):
        cell.data_type = "s"
        cell.quotePrefix = True


def collisions(rows: list[dict[str, Any]]) -> list[str]:
    seen: dict[tuple[str, str], int] = {}
    for r in rows:
        key = (norm(r["system"]), norm(r["terminal"]))
        seen[key] = seen.get(key, 0) + 1
    return [f"{s or '(no system)'} / {t}" for (s, t), n in sorted(seen.items()) if n > 1]


def rows_of(terminals: list[dict[str, Any]], space_names: dict[str, str]) -> list[dict[str, Any]]:
    """Terminals with an airflow, one row per unit (a quantity of 3 is three rows T1-1 .. T1-3), in a stable order."""
    out = []
    for t in sorted(terminals, key=lambda r: (r.get("system_tag") or "", r["tag"])):
        if t.get("airflow_ls") is None:
            continue
        n = int(t.get("quantity") or 1)
        for i in range(n):
            out.append({"system": t.get("system_tag") or "", "terminal": t["tag"] if n == 1 else f"{t['tag']}-{i + 1}", "space": space_names.get(str(t.get("space_id")), ""),
                        "design_ls": float(t["airflow_ls"])})
    return out


def to_xlsx(rows: list[dict[str, Any]], revision_id: str) -> bytes:
    from openpyxl import Workbook

    wb = Workbook()
    info = wb.active
    info.title = "Read me"
    for line in (BANNER, f"Revision {revision_id}", "Fill in the measured airflow (L/s), date, name and comment. Do not change the other columns: the import checks them against the signed design.",
                 "The tolerance is given when the sheet is imported."):
        info.append([line])
    ws = wb.create_sheet(SHEET)
    ws.append(HEADER)
    for n, r in enumerate(rows, start=2):
        ws.append([None, None, None, r["design_ls"], None, None, None, None])
        for col, key in ((1, "system"), (2, "terminal"), (3, "space")):
            put_text(ws.cell(n, col), r[key])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_pdf(rows: list[dict[str, Any]], revision_id: str) -> bytes:
    buf = io.BytesIO()

    def decorate(canvas: Any, doc: Any) -> None:
        canvas.saveState()
        canvas.setFillColor(colors.HexColor("#B00020"))
        canvas.setFont("Helvetica-Bold", 9)
        canvas.drawCentredString(landscape(A4)[0] / 2, landscape(A4)[1] - 9 * mm, BANNER)
        canvas.setFillColor(colors.black)
        canvas.setFont("Helvetica", 7)
        canvas.drawString(12 * mm, 8 * mm, f"Revision {revision_id} | page {doc.page}")
        canvas.restoreState()

    styles = getSampleStyleSheet()
    cell = ParagraphStyle("cell", parent=styles["BodyText"], fontSize=8, leading=10)
    doc = SimpleDocTemplate(buf, pagesize=landscape(A4), topMargin=18 * mm, bottomMargin=14 * mm, leftMargin=12 * mm, rightMargin=12 * mm, title="Commissioning sheet", invariant=1)
    story: list[Any] = []
    systems = sorted({r["system"] for r in rows}) or [""]
    for n, system in enumerate(systems):
        if n:
            story.append(PageBreak())
        story += [Paragraph(escape(f"Commissioning sheet: system {system or '(none)'}"), styles["Title"]), Spacer(1, 2 * mm)]
        data = [[Paragraph(h, cell) for h in ("Terminal", "Space", "Design L/s", "Measured L/s", "Variance %", "Date", "By", "Comment")]]
        for r in [x for x in rows if x["system"] == system]:
            data.append([Paragraph(escape(clean(r["terminal"])), cell), Paragraph(escape(clean(r["space"])), cell), f"{r['design_ls']:g}", "", "", "", "", ""])
        t = Table(data, colWidths=[30 * mm, 50 * mm, 22 * mm, 28 * mm, 22 * mm, 28 * mm, 30 * mm, 60 * mm], repeatRows=1, rowHeights=[7 * mm] + [9 * mm] * (len(data) - 1))
        t.setStyle(TableStyle([("GRID", (0, 0), (-1, -1), 0.4, colors.grey), ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEEEEE")), ("VALIGN", (0, 0), (-1, -1), "MIDDLE")]))
        story.append(t)
    doc.build(story, onFirstPage=decorate, onLaterPages=decorate)
    return buf.getvalue()


def parse_import(data: bytes, rows: list[dict[str, Any]], tolerance_pct: float) -> list[dict[str, Any]]:
    """Readings from a filled-in sheet, matched to the signed design rows. Raises ImportRefused for a file that is not our sheet."""
    from openpyxl import load_workbook

    if len(data) > MAX_IMPORT_BYTES:
        raise ImportRefused("the file is larger than 5 MiB")
    if not math.isfinite(tolerance_pct) or not 0 <= tolerance_pct <= 50:
        raise ImportRefused("the tolerance is a percentage between 0 and 50")
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:                               # a small file must not unpack to something huge
            infos = z.infolist()
            if len(infos) > 200 or sum(i.file_size for i in infos) > MAX_UNPACKED or any(i.file_size > 1000 * max(i.compress_size, 1) and i.file_size > 1_000_000 for i in infos):
                raise ImportRefused("the workbook unpacks to far more than its size (refused)")
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb[SHEET]
        grid = list(ws.iter_rows(min_row=1, max_row=MAX_ROWS + 2, max_col=len(HEADER), values_only=True))
    except ImportRefused:
        raise
    except Exception:  # noqa: BLE001 - any unreadable workbook is refused, not a crash
        raise ImportRefused("that is not a commissioning sheet from this app (no readable 'Terminals' sheet)") from None
    if not grid or [clean(c) for c in grid[0][: len(HEADER)]] != HEADER:
        raise ImportRefused("the header row was changed: use the sheet as downloaded")
    if len(grid) > MAX_ROWS + 1:
        raise ImportRefused(f"more than {MAX_ROWS} rows")
    if collisions(rows):
        raise ImportRefused("terminal tags collide within a system: " + ", ".join(collisions(rows)[:5]))
    design = {(norm(r["system"]), norm(r["terminal"])): r for r in rows}
    seen: set[tuple[str, str]] = set()
    out: list[dict[str, Any]] = []
    for g in grid[1:]:
        if not g or all(c is None or c == "" for c in g):
            continue
        tag, system = norm(g[1]), norm(g[0])
        key = (system, tag)
        d = design.get(key)
        base = {"terminal": tag, "system": system, "design_ls": None if d is None else d["design_ls"], "measured_ls": None, "variance_pct": None,
                "tolerance_pct": tolerance_pct, "measured_on": clean(g[5])[:40], "measured_by": clean(g[6])[:80], "comment": clean(g[7])[:300]}
        if d is None:
            out.append({**base, "flag": "NOT_IN_DESIGN"})
            continue
        if key in seen:
            out.append({**base, "flag": "DUPLICATE"})
            continue
        seen.add(key)
        try:
            sheet_design = float(g[3])
        except (TypeError, ValueError, OverflowError):
            sheet_design = math.nan
        if not math.isclose(sheet_design, d["design_ls"], rel_tol=1e-9, abs_tol=1e-9):
            out.append({**base, "flag": "DESIGN_CHANGED"})
            continue
        m = g[COL_MEASURED]
        if m is None or m == "":
            out.append({**base, "flag": "NOT_MEASURED"})
            continue
        try:
            mv = float(m) if isinstance(m, int | float) and not isinstance(m, bool) else math.nan
        except OverflowError:
            mv = math.nan
        if not math.isfinite(mv) or mv < 0:
            out.append({**base, "flag": "INVALID"})
            continue
        var = (mv - d["design_ls"]) / d["design_ls"] * 100.0 if d["design_ls"] else math.nan
        flag = "INVALID" if not math.isfinite(var) else ("WITHIN_TOLERANCE" if abs(var) <= tolerance_pct else "OUTSIDE_TOLERANCE")
        out.append({**base, "measured_ls": mv, "variance_pct": None if not math.isfinite(var) else round(var, 2), "flag": flag})
    for key in sorted(set(design) - seen):
        out.append({"terminal": key[1], "system": design[key]["system"], "design_ls": design[key]["design_ls"], "measured_ls": None, "variance_pct": None, "tolerance_pct": tolerance_pct,
                    "measured_on": "", "measured_by": "", "comment": "", "flag": "MISSING_FROM_SHEET"})
    return out
