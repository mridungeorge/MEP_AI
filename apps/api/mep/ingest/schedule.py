"""System/equipment schedule: the published Excel template and its reader.

The template has one column per rule input of ONE NCC edition (read from the rule YAML, never hard-coded), with the
unit in the header, e.g. ``control_deadband [K]``. The reader is machine reading: every value it returns carries
provenance 'extracted' (source 'xlsx') and the engine refuses it until a designer confirms it at Gate 1. Nothing here
decides compliance.

Units go through mep.engine.units (strict grammar, declared-dimension check). A compatible unit (m^3/h for L/s) is
converted to the rule's declared unit; an incompatible one rejects the row. Formulas, macros and external links are
refused; hidden sheets are ignored.
"""
import hashlib
import io
import math
import re
import zipfile
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl.worksheet.datavalidation import DataValidation

from mep.engine.loader import InputSpec, RulePack
from mep.engine.units import UnitError, coerce

TEMPLATE_VERSION = "1"
MAX_ROWS = 2000
MAX_BYTES = 20 * 1024 * 1024
MAX_UNCOMPRESSED = 200 * 1024 * 1024
MAX_SCANNED_ROWS = 20000
MAX_ERRORS = 500
MAX_TEXT = 500
SYSTEMS_SHEET = "Systems"
TAG_HEADER = "system tag"
TYPE_HEADER = "system type"
DIMENSIONLESS = "dimensionless"
NAME_RE = re.compile(r"^[a-z][a-z0-9_]{0,63}$")
_HEADER = re.compile(r"^(?P<name>[A-Za-z][A-Za-z0-9_]*)\s*(?:\[(?P<unit>[^\]]*)\])?$")
_EDITION = re.compile(r"^NCC(2022|2025)$")
_FIXED_TIME = datetime(2026, 1, 1)  # noqa: DTZ001 - xlsx core properties are naive


@dataclass(frozen=True)
class ScheduleInput:
    name: str
    value: float | int | str | bool
    unit: str | None
    provenance: str = "extracted"
    source: str = "xlsx"


@dataclass(frozen=True)
class ScheduleRow:
    tag: str
    system_type: str
    inputs: list[ScheduleInput]
    row: int = 0


@dataclass(frozen=True)
class ScheduleError:
    row: int | None
    column: str | None
    message: str


@dataclass
class ScheduleResult:
    rows: list[ScheduleRow] = field(default_factory=list)
    errors: list[ScheduleError] = field(default_factory=list)


class CellError(ValueError):
    """A value that cannot be used for the declared input."""


# --- rule inputs -------------------------------------------------------------------------------------------------

PROJECT_FACTS = ("climate_zone", "building_class")   # come from the project (building parts), never from a schedule row


def schedule_inputs(pack: RulePack, edition: str) -> dict[str, InputSpec]:
    """The inputs declared by the rules of one edition, by name. Editions are never mixed."""
    if not _EDITION.match(edition):
        raise ValueError(f"unknown NCC edition {edition!r}")
    found: dict[str, InputSpec] = {}
    for rule in pack.rules.values():
        if rule.edition != edition:
            continue
        for name, spec in rule.inputs.items():
            if name in PROJECT_FACTS:
                continue
            old = found.get(name)
            if old is not None and (old.type, old.unit, old.values) != (spec.type, spec.unit, spec.values):
                raise ValueError(f"input {name!r} is declared differently by different {edition} rules")
            found[name] = spec
    if not found:
        raise ValueError(f"no rules for edition {edition!r}")
    return dict(sorted(found.items()))


def _used_by(pack: RulePack, edition: str) -> dict[str, list[str]]:
    used: dict[str, list[str]] = {}
    for rule in pack.rules.values():
        if rule.edition == edition:
            for name in rule.inputs:
                used.setdefault(name, []).append(rule.id)
    return {k: sorted(v) for k, v in used.items()}


def system_types(pack: RulePack, edition: str) -> tuple[str, ...] | None:
    """Allowed system types: the rules' own `system_type` enumeration, or None when the edition declares none."""
    spec = schedule_inputs(pack, edition).get("system_type")
    return spec.values if spec is not None and spec.values else None


def declared_unit(spec: InputSpec) -> str | None:
    """The unit of a numeric input (dimensionless when the rule declares none); None for non-numbers."""
    if spec.type in ("number", "int"):
        return spec.unit or DIMENSIONLESS
    return None


def column_header(name: str, spec: InputSpec) -> str:
    unit = declared_unit(spec)
    return f"{name} [{unit}]" if unit else name


# --- value normalisation (shared with the API form) --------------------------------------------------------------

def normalise_value(spec: InputSpec, value: Any, unit: str | None) -> tuple[float | int | str | bool, str | None]:
    """Check one value against its rule input and return it in the rule's declared unit. Raises CellError."""
    if spec.type in ("number", "int"):
        target = declared_unit(spec) or DIMENSIONLESS
        try:
            quantity = coerce(value, target, unit)
            number = float(value) if unit == target else float(quantity.to(target).magnitude)
        except UnitError as exc:
            raise CellError(str(exc)) from exc
        number = float(f"{number:.12g}")
        if spec.type == "int":
            if not number.is_integer():
                raise CellError("must be a whole number")
            return int(number), target
        return number, target
    if unit not in (None, ""):
        raise CellError("a unit is only for numeric inputs")
    if spec.type == "bool":
        if not isinstance(value, bool):
            raise CellError("must be TRUE or FALSE")
        return value, None
    if not isinstance(value, str) or not value.strip() or len(value) > MAX_TEXT:
        raise CellError("must be text")
    text = value.strip()
    if spec.type == "enum" and (spec.values is None or text not in spec.values):
        raise CellError(f"must be one of: {', '.join(spec.values or ())}")
    return text, None


# --- template ----------------------------------------------------------------------------------------------------

def _digest(inputs: dict[str, InputSpec]) -> str:
    text = "|".join(f"{n}:{s.type}:{s.unit}:{s.values}" for n, s in inputs.items())
    return hashlib.sha256(text.encode()).hexdigest()[:12]


def build_template(pack: RulePack, edition: str) -> bytes:
    """The xlsx template for one edition. Deterministic: the same rules give the same cells."""
    inputs = schedule_inputs(pack, edition)
    types = system_types(pack, edition)
    wb = openpyxl.Workbook()
    systems = wb.active
    systems.title = SYSTEMS_SHEET
    readme = wb.create_sheet("ReadMe")
    lists = wb.create_sheet("Lists")

    systems.append([TAG_HEADER, TYPE_HEADER] + [column_header(n, s) for n, s in inputs.items()])
    systems.freeze_panes = "C2"
    for col in range(1, 3 + len(inputs)):
        systems.column_dimensions[openpyxl.utils.get_column_letter(col)].width = 22
    last = MAX_ROWS + 1

    # Lists: system types, TRUE/FALSE, and one column per enumeration
    list_cols: dict[str, str] = {}
    lists_headers: list[str] = []
    columns: list[tuple[str, list[str]]] = []
    if types:
        columns.append(("system type", list(types)))
    columns.append(("boolean", ["TRUE", "FALSE"]))
    for name, spec in inputs.items():
        if spec.type == "enum" and spec.values:
            columns.append((name, list(spec.values)))
    for idx, (title, values) in enumerate(columns, start=1):
        letter = openpyxl.utils.get_column_letter(idx)
        lists.cell(row=1, column=idx, value=title)
        for r, v in enumerate(values, start=2):
            lists.cell(row=r, column=idx, value=v)
        list_cols[title] = f"=Lists!${letter}$2:${letter}${len(values) + 1}"
        lists_headers.append(title)

    def validate(dv: DataValidation, col: int) -> None:
        letter = openpyxl.utils.get_column_letter(col)
        dv.add(f"{letter}2:{letter}{last}")
        systems.add_data_validation(dv)

    if types:
        validate(DataValidation(type="list", formula1=list_cols["system type"], allow_blank=True), 2)
    for offset, (name, spec) in enumerate(inputs.items(), start=3):
        if spec.type == "bool":
            validate(DataValidation(type="list", formula1=list_cols["boolean"], allow_blank=True), offset)
        elif spec.type == "enum" and spec.values:
            validate(DataValidation(type="list", formula1=list_cols[name], allow_blank=True), offset)
        elif spec.type == "int":
            validate(DataValidation(type="whole", operator="between", formula1="-1000000000",
                                    formula2="1000000000", allow_blank=True), offset)
        elif spec.type == "number":
            validate(DataValidation(type="decimal", operator="between", formula1="-1E+15", formula2="1E+15",
                                    allow_blank=True), offset)

    readme.append(["MEP system schedule template"])
    readme.append(["Template version", TEMPLATE_VERSION])
    readme.append(["NCC edition", edition])
    readme.append(["Input set digest", _digest(inputs)])
    readme.append(["Inputs", len(inputs)])
    readme.append(["How to use", ("One row per system on the Systems sheet. Keep the header text as it is: the unit"
                   " in square brackets is the unit of the column. A compatible unit (for example m^3/h in place"
                   " of L/s) is converted; a unit of another quantity rejects the row.")])
    readme.append(["Rules of the file", ("No formulas, macros or external links (values only). Hidden sheets are"
                   f" ignored. At most {MAX_ROWS} systems. Only .xlsx.")])
    readme.append(["Status of values", ("Values read from this file are 'extracted'. A designer must confirm them"
                   " at Gate 1 before any rule uses them.")])
    readme.append([])
    readme.append(["Input", "Type", "Unit", "Allowed values", "Used by rules"])
    for name, spec in inputs.items():
        readme.append([name, spec.type, declared_unit(spec) or "", ", ".join(spec.values or ()),
                       ", ".join(_used_by(pack, edition)[name])])
    readme.column_dimensions["A"].width = 38
    readme.column_dimensions["B"].width = 30

    wb.properties.creator = "mep-copilot"
    wb.properties.created = _FIXED_TIME
    wb.properties.modified = _FIXED_TIME
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# --- reader ------------------------------------------------------------------------------------------------------

def _check_package(data: bytes) -> str | None:
    if not data:
        return "the file is empty"
    if len(data) > MAX_BYTES:
        return "the file is too large"
    if not zipfile.is_zipfile(io.BytesIO(data)):
        return "not an .xlsx workbook"
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            names = z.namelist()
            if sum(i.file_size for i in z.infolist()) > MAX_UNCOMPRESSED:
                return "the workbook expands to an unreasonable size"
            if "xl/workbook.xml" not in names or "[Content_Types].xml" not in names:
                return "not an .xlsx workbook"
            lowered = [n.lower() for n in names]
            if any("vbaproject" in n or n.endswith(".bin") and "macro" in n for n in lowered):
                return "workbooks with macros are refused (use plain .xlsx)"
            if any(n.startswith("xl/externallinks/") for n in lowered):
                return "workbooks with external links are refused"
            if b"macroEnabled" in z.read("[Content_Types].xml"):
                return "macro-enabled workbooks are refused (use plain .xlsx)"
    except (zipfile.BadZipFile, KeyError, OSError):
        return "not an .xlsx workbook"
    return None


def read_schedule(source: bytes | bytearray | str | Path, pack: RulePack, edition: str) -> ScheduleResult:
    """Read the Systems sheet. Rows with any problem are rejected whole; header problems are reported."""
    result = ScheduleResult()
    inputs = schedule_inputs(pack, edition)
    types = system_types(pack, edition)

    def fail(message: str, row: int | None = None, column: str | None = None) -> None:
        if len(result.errors) < MAX_ERRORS:
            result.errors.append(ScheduleError(row, column, message))

    if isinstance(source, str | Path):
        path = Path(source)
        if path.suffix.lower() != ".xlsx":
            fail("only .xlsx files are accepted")
            return result
        try:
            if path.stat().st_size > MAX_BYTES:
                fail("the file is too large")
                return result
            data = path.read_bytes()
        except OSError as exc:
            fail(f"cannot read the file: {exc.strerror or exc}")
            return result
    else:
        data = bytes(source)
    problem = _check_package(data)
    if problem:
        fail(problem)
        return result
    try:
        wb = openpyxl.load_workbook(io.BytesIO(data), data_only=False, keep_vba=False, keep_links=False)
    except Exception:  # noqa: BLE001 - any parser failure is just an unreadable file
        fail("the workbook cannot be read")
        return result
    sheet = next((ws for ws in wb.worksheets if ws.title.strip().lower() == SYSTEMS_SHEET.lower()), None)
    if sheet is None or sheet.sheet_state != "visible":
        fail(f"no visible sheet named {SYSTEMS_SHEET!r}")
        return result

    last_row = min(sheet.max_row, MAX_SCANNED_ROWS + 1)
    rows_iter = sheet.iter_rows(min_row=1, max_row=last_row)
    header_cells = next(rows_iter, ())
    columns: list[tuple[str, str | None, InputSpec | None] | None] = []  # per column: (name, unit, spec)
    tag_col = type_col = None
    seen: set[str] = set()
    for idx, cell in enumerate(header_cells):
        raw = cell.value
        if raw is None or (isinstance(raw, str) and not raw.strip()):
            columns.append(None)
            continue
        if cell.data_type == "f" or not isinstance(raw, str):
            fail("header cells must be plain text", 1, str(raw))
            columns.append(None)
            continue
        text = " ".join(raw.split())
        if text.lower() == TAG_HEADER:
            tag_col = idx
            columns.append(None)
            continue
        if text.lower() == TYPE_HEADER:
            type_col = idx
            columns.append(None)
            continue
        m = _HEADER.match(text)
        name = m["name"] if m else text
        unit = (m["unit"] or "").strip() or None if m else None
        if not m or name not in inputs:
            fail(f"unknown column {text!r} for {edition}; it is ignored", 1, name)
            columns.append(None)
            continue
        if name in seen:
            fail(f"duplicate column {name!r}", 1, name)
            columns.append(None)
            continue
        seen.add(name)
        columns.append((name, unit, inputs[name]))
    if tag_col is None or type_col is None:
        fail(f"the sheet needs the columns {TAG_HEADER!r} and {TYPE_HEADER!r}", 1)
        return result

    tags: set[str] = set()
    count = 0
    scanned = 0
    for excel_row, cells in enumerate(rows_iter, start=2):
        scanned += 1
        if all(c.value is None or (isinstance(c.value, str) and not c.value.strip()) for c in cells):
            continue
        count += 1
        if count > MAX_ROWS:
            result.rows.clear()
            fail(f"more than {MAX_ROWS} systems; split the schedule", excel_row)
            return result
        errors_before = len(result.errors)
        row_errors: list[ScheduleError] = []

        def bad(message: str, column: str | None = None, _r: int = excel_row,
                _e: list[ScheduleError] = row_errors) -> None:
            _e.append(ScheduleError(_r, column, message))

        def cell_at(idx: int, _cells: Any = cells) -> Any:
            return _cells[idx] if idx < len(_cells) else None

        tag_cell, type_cell = cell_at(tag_col), cell_at(type_col)
        tag = system_type = ""
        if tag_cell is not None and tag_cell.data_type == "f" or type_cell is not None and type_cell.data_type == "f":
            bad("formulas are refused; enter values", TAG_HEADER)
        else:
            tag = str(tag_cell.value).strip() if tag_cell is not None and tag_cell.value is not None else ""
            system_type = str(type_cell.value).strip() if type_cell is not None and type_cell.value is not None else ""
            if not tag or len(tag) > 64:
                bad("system tag is required (64 characters at most)", TAG_HEADER)
            elif tag in tags:
                bad(f"duplicate system tag {tag!r}", TAG_HEADER)
            if not system_type or len(system_type) > 64 or (types is not None and system_type not in types):
                allowed = f" (one of: {', '.join(types)})" if types else ""
                bad(f"system type is required{allowed}", TYPE_HEADER)
        found: list[ScheduleInput] = []
        for idx, col in enumerate(columns):
            cell = cell_at(idx)
            if col is None or cell is None or cell.value is None:
                continue
            name, unit, spec = col
            assert spec is not None
            if isinstance(cell.value, str) and not cell.value.strip():
                continue
            if cell.data_type == "f":
                bad("formulas are refused; enter values", name)
                continue
            value = cell.value
            if spec.type in ("number", "int"):
                if isinstance(value, bool) or not isinstance(value, int | float):
                    bad("text or other non-number in a number cell", name)
                    continue
                if not math.isfinite(value):
                    bad("not a finite number", name)
                    continue
                if unit is None:
                    bad(f"the column header has no unit; write it as {name} [{declared_unit(spec)}]", name)
                    continue
            elif isinstance(value, str):
                value = value.strip()
                if spec.type == "bool" and value.upper() in ("TRUE", "FALSE"):
                    value = value.upper() == "TRUE"
            try:
                norm, norm_unit = normalise_value(spec, value, unit if spec.type in ("number", "int") else None)
            except CellError as exc:
                bad(f"{name}: {exc}", name)
                continue
            found.append(ScheduleInput(name, norm, norm_unit))
        if row_errors:
            for e in row_errors:
                fail(e.message, e.row, e.column)
            continue
        assert len(result.errors) == errors_before
        tags.add(tag)
        result.rows.append(ScheduleRow(tag, system_type, found, excel_row))
    if sheet.max_row > MAX_SCANNED_ROWS + 1:
        fail(f"the sheet runs past row {MAX_SCANNED_ROWS}; trailing rows were not read", MAX_SCANNED_ROWS)
    return result
